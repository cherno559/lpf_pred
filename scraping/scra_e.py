import cloudscraper
import pandas as pd
import matplotlib.pyplot as plt
from mplsoccer import Pitch
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font
from openpyxl.formatting.rule import ColorScaleRule
import os
import traceback
import time
from datetime import datetime

# =========================================================
# 1. CONFIGURACIÓN DINÁMICA DE TEMPORADA
# =========================================================
EQUIPO_OBJETIVO = "River Plate" 
TEAM_ID = 3211  # ID de River en SofaScore
ANIO_TEMPORADA = 2026

# --- RUTA PARA LOS DATOS INDIVIDUALES ---
CARPETA_TRABAJO = os.path.join("data", "equipos")
os.makedirs(CARPETA_TRABAJO, exist_ok=True) 

RUTA_EXCEL = os.path.join(CARPETA_TRABAJO, f"Base_Datos_{EQUIPO_OBJETIVO.replace(' ', '_')}_{ANIO_TEMPORADA}.xlsx")

# =========================================================
# BYPASS DEFINITIVO: CLOUDSCRAPER
# =========================================================
session = cloudscraper.create_scraper(
    browser={
        'browser': 'chrome',
        'platform': 'windows',
        'desktop': True
    }
)

headers = {
    "Accept": "application/json, text/plain, */*",
    "Accept-Language": "es-ES,es;q=0.9,en;q=0.8",
    "Cache-Control": "no-cache",
    "Origin": "https://www.sofascore.com",
    "Referer": "https://www.sofascore.com/",
    "Sec-Ch-Ua": '"Chromium";v="122", "Not(A:Brand";v="24", "Google Chrome";v="122"',
    "Sec-Ch-Ua-Mobile": "?0",
    "Sec-Ch-Ua-Platform": '"Windows"',
    "Sec-Fetch-Dest": "empty",
    "Sec-Fetch-Mode": "cors",
    "Sec-Fetch-Site": "same-site",
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
}

ESTADISTICAS_CLAVE = {
    "Ball possession": "Posesión de balón",
    "Expected goals": "Goles esperados (xG)",
    "Total shots": "Tiros totales",
    "Shots on target": "Tiros al arco",
    "Shots off target": "Tiros afuera",
    "Blocked shots": "Tiros bloqueados",
    "Corner kicks": "Córners",
    "Offsides": "Fueras de juego",
    "Fouls": "Faltas",
    "Yellow cards": "Tarjetas amarillas",
    "Red cards": "Tarjetas rojas",
    "Passes": "Pases totales",
    "Accurate passes": "Pases precisos",
    "Accurate long balls": "Balones largos precisos",
    "Accurate crosses": "Centros precisos",
    "Goalkeeper saves": "Atajadas del arquero",
    "Tackles": "Quites",
    "Interceptions": "Intercepciones",
    "Clearances": "Despejes"
}

def formato_fraccion(acertados, totales): return f"'{acertados}/{totales}"
def calcular_porcentaje(acertados, totales): return int((acertados / totales) * 100) if totales > 0 else 0

def extraer_numero(valor):
    if isinstance(valor, str):
        valor = valor.replace('%', '')
    try: return float(valor)
    except: return 0.0

def obtener_partidos_liga(team_id, anio_temporada):
    print(f"🔍 Buscando partidos de liga para el ID {team_id} en el año {anio_temporada}...")
    partidos_validos = []
    page = 0
    
    while True:
        url = f"https://api.sofascore.com/api/v1/team/{team_id}/events/last/{page}"
        response = session.get(url, headers=headers)
        
        if response.status_code != 200:
            print(f"⚠️ Error {response.status_code} al consultar el historial.")
            print(f"🔍 Contenido recibido: {response.text[:300]}")
            break
            
        data = response.json()
        if 'events' not in data or not data['events']:
            break
            
        for event in data['events']:
            if event.get('status', {}).get('type') != 'finished': 
                continue
                
            ts = event.get('startTimestamp')
            if not ts: continue
            dt = datetime.fromtimestamp(ts)
            
            if dt.year != anio_temporada: 
                continue
                
            torneo = event.get('tournament', {}).get('name', '').upper()
            if 'LIGA PROFESIONAL' in torneo or 'COPA DE LA LIGA' in torneo:
                partidos_validos.append({
                    "id": event['id'],
                    "fecha": dt.strftime('%d-%m'),
                    "torneo": torneo
                })
        
        page += 1
        time.sleep(1.5)
        
        if dt.year < anio_temporada:
            break
            
    print(f"✅ Se encontraron {len(partidos_validos)} partidos de Liga/Copa local en {anio_temporada}.")
    return partidos_validos

def ejecutar_reporte_exacto(event_id, equipo_objetivo):
    print(f"\n--- Generando Reporte (ID: {event_id}) ---")

    try:
        response = session.get(f"https://api.sofascore.com/api/v1/event/{event_id}", headers=headers)
        if response.status_code != 200:
            print(f"❌ Error HTTP {response.status_code}.")
            print(f"🔍 Contenido recibido: {response.text[:300]}")
            return
            
        res_info = response.json()
        if 'event' not in res_info:
            print("❌ Error: Datos no encontrados.")
            return
            
        event_data = res_info.get('event', {})
        home_data = event_data.get('homeTeam') or event_data.get('home_team')
        away_data = event_data.get('awayTeam') or event_data.get('away_team')
        
        local_name = home_data['name']
        visitante_name = away_data['name']
        
        if equipo_objetivo.lower() in local_name.lower():
            lado_equipo = 'home'
            rival = visitante_name
        elif equipo_objetivo.lower() in visitante_name.lower():
            lado_equipo = 'away'
            rival = local_name
        else:
            print(f"❌ Error: El equipo '{equipo_objetivo}' no está en este partido.")
            return

        ts = event_data.get('startTimestamp')
        fecha_partido_dt = datetime.fromtimestamp(ts) if ts else datetime.now()
        fecha_partido_str = fecha_partido_dt.strftime('%d-%m')
        condicion = "(L)" if lado_equipo == 'home' else "(V)"
        
        torneo_nombre = event_data.get('tournament', {}).get('name', '').upper()
        if 'COPA DE LA LIGA' in torneo_nombre: comp_abrev = 'CDL'
        elif 'LIGA PROFESIONAL' in torneo_nombre or 'LIGA' in torneo_nombre: comp_abrev = 'LPF'
        else: comp_abrev = torneo_nombre[:3]

        ronda_info = event_data.get('roundInfo', {})
        ronda_num = ronda_info.get('round')
        f_str = f"F{ronda_num}" if ronda_num else "P"
        
        prefijo = f"{fecha_partido_str} {f_str} VS "
        sufijo = f" {condicion} {comp_abrev}"
        caracteres_disp = 31 - len(prefijo) - len(sufijo)
        rival_corto = rival.upper()[:max(3, caracteres_disp)].strip()
        nombre_hoja = f"{prefijo}{rival_corto}{sufijo}"[:31]

        goles_local = event_data.get('homeScore', {}).get('current', 0)
        goles_visitante = event_data.get('awayScore', {}).get('current', 0)

        res_stats = session.get(f"https://api.sofascore.com/api/v1/event/{event_id}/statistics", headers=headers).json()
        
        stats_temporales = {}
        if 'statistics' in res_stats and len(res_stats['statistics']) > 0:
            stats_all = next((p for p in res_stats['statistics'] if p['period'] == 'ALL'), None)
            if stats_all:
                for grupo in stats_all['groups']:
                    for item in grupo['statisticsItems']:
                        nombre_original = item.get('name', '')
                        if nombre_original not in stats_temporales:
                            stats_temporales[nombre_original] = {
                                'home': item.get('home', '0'),
                                'away': item.get('away', '0')
                            }

        estadisticas_partido = []
        stats_num_equipo = {}
        stats_num_rival = {}

        for clave_ingles, nombre_espanol in ESTADISTICAS_CLAVE.items():
            if clave_ingles in stats_temporales:
                val_home = stats_temporales[clave_ingles]['home']
                val_away = stats_temporales[clave_ingles]['away']
                estadisticas_partido.append([nombre_espanol, val_home, val_away])
                
                if lado_equipo == 'home':
                    stats_num_equipo[nombre_espanol] = extraer_numero(val_home)
                    stats_num_rival[nombre_espanol] = extraer_numero(val_away)
                else:
                    stats_num_equipo[nombre_espanol] = extraer_numero(val_away)
                    stats_num_rival[nombre_espanol] = extraer_numero(val_home)

        if not estadisticas_partido:
            print("⚠️️ No hay estadísticas detalladas. Se salta el partido.")
            return

        estadisticas_partido.insert(0, ["Resultado", goles_local, goles_visitante])

        puntos_altos = []
        puntos_bajos = []

        if stats_num_equipo.get("Posesión de balón", 0) > 55: puntos_altos.append("Dominio de la posesión de balón")
        elif stats_num_equipo.get("Posesión de balón", 0) < 45: puntos_bajos.append("Dificultad para retener la pelota")
        if stats_num_equipo.get("Tiros al arco", 0) >= 5: puntos_altos.append("Buen volumen de tiros al arco")
        elif stats_num_equipo.get("Tiros totales", 0) < 8: puntos_bajos.append("Poca generación de juego ofensivo")
        if stats_num_rival.get("Tiros al arco", 0) <= 2: puntos_altos.append("Solidez defensiva (Permitió pocos tiros)")
        elif stats_num_rival.get("Tiros al arco", 0) >= 5: puntos_bajos.append("El rival llegó con facilidad al arco")
        if stats_num_equipo.get("Tiros afuera", 0) > stats_num_equipo.get("Tiros al arco", 0) and stats_num_equipo.get("Tiros totales", 0) > 10:
            puntos_bajos.append("Baja efectividad en la definición")

        if not puntos_altos: puntos_altos.append("Sin aspectos destacados")
        if not puntos_bajos: puntos_bajos.append("Sin aspectos negativos evidentes")

        # JUGADORES
        res_lineups = session.get(f"https://api.sofascore.com/api/v1/event/{event_id}/lineups", headers=headers).json()
        jugadores_equipo = res_lineups[lado_equipo]['players']
        
        orden_pos = {'G': (1, 'Arquero'), 'D': (2, 'Defensor'), 'M': (3, 'Mediocampista'), 'F': (4, 'Delantero')}
        filas_excel, dicc_plantel = [], {}

        for j in jugadores_equipo:
            s = j.get('statistics', {})
            minutos = s.get('minutesPlayed', 0)
            if minutos == 0: continue
            
            p_id = j['player']['id']
            dicc_plantel[p_id] = {
                'nombre': j['player'].get('shortName', '').split()[-1].upper(),
                'numero': j.get('shirtNumber', ''),
                'es_titular': not j.get('substitute', False)
            }

            p_acc, p_tot = s.get('accuratePass', 0), s.get('totalPass', 0)
            d_gan, d_tot = s.get('duelWon', 0), s.get('duelWon', 0) + s.get('duelLost', 0)
            c_acc, c_tot = s.get('accurateCross', 0), s.get('totalCross', 0)
            b_acc, b_tot = s.get('accurateLongBalls', 0), s.get('totalLongBalls', 0)
            r_acc, r_tot = s.get('wonContest', 0), s.get('totalContest', 0)
            t_arco = s.get('onTargetScoringAttempt', 0)
            t_afuera = s.get('shotOffTarget', 0)
            t_tot = t_arco + t_afuera + s.get('blockedScoringAttempt', 0)

            filas_excel.append({
                'Jugador': j['player']['name'], 
                'Posición': orden_pos.get(j['player'].get('position', 'M'), (5, 'Otro'))[1],
                'Orden': orden_pos.get(j['player'].get('position', 'M'), (5, 'Otro'))[0],
                'Minutos': minutos, 
                'Nota SofaScore': s.get('rating', 0.0),
                'Quites (Tackles)': s.get('totalTackle', 0), 
                'Intercepciones': s.get('interceptionWon', 0),
                'Despejes': s.get('totalClearance', 0),
                'Duelos (Gan/Tot)': formato_fraccion(d_gan, d_tot), 
                'Efectividad Duelos': calcular_porcentaje(d_gan, d_tot),
                'Pérdidas de posesión': s.get('possessionLostCtrl', 0),
                'Faltas Cometidas': s.get('fouls', 0), 
                'Faltas Recibidas': s.get('wasFouled', 0),
                'Asistencias': s.get('goalAssist', 0), 
                'Pases Clave': s.get('keyPass', 0),
                'Pases (Comp/Tot)': formato_fraccion(p_acc, p_tot), 
                'Efectividad Pases': calcular_porcentaje(p_acc, p_tot),
                'Centros (Comp/Tot)': formato_fraccion(c_acc, c_tot), 
                'Efectividad Centros': calcular_porcentaje(c_acc, c_tot),
                'Balones Largos (C/T)': formato_fraccion(b_acc, b_tot), 
                'Efectividad Balones': calcular_porcentaje(b_acc, b_tot),
                'Regates (Exit/Tot)': formato_fraccion(r_acc, r_tot), 
                'Efectividad Regates': calcular_porcentaje(r_acc, r_tot),
                'Goles': s.get('goals', 0), 
                'Tiros Totales': t_tot, 
                'Tiros al Arco': t_arco, 
                'Tiros Afuera': t_afuera
            })

        df = pd.DataFrame(filas_excel).sort_values(by=['Orden', 'Minutos'], ascending=[True, False]).drop(columns=['Orden'])

        # GRAFICOS
        pitch = Pitch(pitch_type='opta', pitch_color='#1a1a1a', line_color='#555555')
        fig1, ax1 = pitch.draw(figsize=(8, 5.5))
        fig1.patch.set_facecolor('#1a1a1a')
        
        try:
            res_pos = session.get(f"https://api.sofascore.com/api/v1/event/{event_id}/average-positions", headers=headers).json()
            for p in res_pos.get(lado_equipo, []):
                if 'player' not in p: continue
                id_j = p['player']['id']
                if id_j in dicc_plantel and dicc_plantel[id_j]['es_titular']:
                    jug = dicc_plantel[id_j]
                    ax1.scatter(p['averageX'], p['averageY'], s=500, color="#d32f2f", edgecolor="white", zorder=5)
                    ax1.text(p['averageX'], p['averageY'], str(jug['numero']), color="white", fontsize=9, ha="center", va="center", fontweight="bold")
                    ax1.text(p['averageX'], p['averageY'] + 4.5, jug['nombre'], color="white", fontsize=7, ha="center", fontweight='bold')
        except:
            pass

        plt.title(f"POSICIONES MEDIAS - {equipo_objetivo.upper()}", color="white", fontsize=12, pad=10)
        img_parado = os.path.join(CARPETA_TRABAJO, f"parado_{event_id}.png")
        plt.savefig(img_parado, bbox_inches="tight", dpi=120, facecolor='#1a1a1a')
        plt.close(fig1)

        fig2, ax2 = pitch.draw(figsize=(8, 5.5))
        fig2.patch.set_facecolor('#1a1a1a')
        fig3, ax3 = pitch.draw(figsize=(8, 5.5))
        fig3.patch.set_facecolor('#1a1a1a')
        
        try:
            res_shots = session.get(f"https://api.sofascore.com/api/v1/event/{event_id}/shotmap", headers=headers).json()
            for s in res_shots.get('shotmap', []):
                x, y = s['playerCoordinates']['x'], s['playerCoordinates']['y']
                nombre = s['player']['shortName'].split()[-1].upper()
                tipo = s['shotType']
                
                if tipo == 'goal': color, marker, size = "#FFD700", "*", 250 
                elif tipo == 'save': color, marker, size = "#3498db", "s", 100 
                elif tipo == 'block': color, marker, size = "#9b59b6", "^", 100 
                else: color, marker, size = "#e74c3c", "o", 100 
                
                if s['isHome'] == (lado_equipo == 'home'):
                    ax2.scatter(x, y, s=size, color=color, marker=marker, edgecolor="white", alpha=0.9, zorder=5)
                    ax2.text(x, y + 3, nombre, color="white", fontsize=6, ha="center")
                else:
                    ax3.scatter(x, y, s=size, color=color, marker=marker, edgecolor="white", alpha=0.9, zorder=5)
                    ax3.text(x, y + 3, nombre, color="white", fontsize=6, ha="center")
        except:
            pass

        for ax in [ax2, ax3]:
            ax.scatter(5, 95, s=100, color="#FFD700", marker="*", edgecolor="white")
            ax.text(7, 94.5, "Gol", color="white", fontsize=8)
            ax.scatter(18, 95, s=60, color="#3498db", marker="s", edgecolor="white")
            ax.text(20, 94.5, "Atajado", color="white", fontsize=8)

        ax2.set_title(f"MAPA DE TIROS - {equipo_objetivo.upper()}", color="white", fontsize=12, pad=10)
        img_shots = os.path.join(CARPETA_TRABAJO, f"shots_{event_id}.png")
        fig2.savefig(img_shots, bbox_inches="tight", dpi=120, facecolor='#1a1a1a')
        plt.close(fig2)

        ax3.set_title(f"MAPA DE TIROS - {rival.upper()}", color="white", fontsize=12, pad=10)
        img_shots_rival = os.path.join(CARPETA_TRABAJO, f"shots_rival_{event_id}.png")
        fig3.savefig(img_shots_rival, bbox_inches="tight", dpi=120, facecolor='#1a1a1a')
        plt.close(fig3)

        # GUARDADO EXCEL
        print(f"Guardando en Excel ({RUTA_EXCEL})...")
        if os.path.exists(RUTA_EXCEL):
            with pd.ExcelWriter(RUTA_EXCEL, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
                df.to_excel(writer, index=False, sheet_name=nombre_hoja)
        else:
            df.to_excel(RUTA_EXCEL, index=False, sheet_name=nombre_hoja)

        wb = load_workbook(RUTA_EXCEL)
        ws = wb[nombre_hoja]
        ws.freeze_panes = 'B2'
        
        for idx in range(1, len(df.columns) + 1):
            ws.column_dimensions[get_column_letter(idx)].width = 16

        rojo, amarillo, verde, blanco = 'F8696B', 'FFEB84', '63BE7B', 'FFFFFF'
        regla_verde_bueno = ColorScaleRule(start_type='min', start_color=rojo, mid_type='percentile', mid_value=50, mid_color=amarillo, end_type='max', end_color=verde)
        regla_rojo_bueno = ColorScaleRule(start_type='min', start_color=verde, mid_type='percentile', mid_value=50, mid_color=amarillo, end_type='max', end_color=rojo)
        regla_solo_verde = ColorScaleRule(start_type='min', start_color=blanco, end_type='max', end_color=verde)
        regla_solo_rojo = ColorScaleRule(start_type='min', start_color=blanco, end_type='max', end_color=rojo)

        columnas_df = list(df.columns)
        ultima_fila = len(df) + 1

        def aplicar_color(nombre_col, regla):
            if nombre_col in columnas_df:
                letra = get_column_letter(columnas_df.index(nombre_col) + 1)
                ws.conditional_formatting.add(f'{letra}2:{letra}{ultima_fila}', regla)

        for col in ['Nota SofaScore', 'Efectividad Duelos', 'Efectividad Pases', 'Efectividad Centros', 'Efectividad Regates']: aplicar_color(col, regla_verde_bueno)
        for col in ['Pérdidas de posesión']: aplicar_color(col, regla_rojo_bueno)
        for col in ['Minutos', 'Quites (Tackles)', 'Intercepciones', 'Despejes', 'Asistencias', 'Goles']: aplicar_color(col, regla_solo_verde)

        fila_inicio_stats = ultima_fila + 3
        ws.cell(row=fila_inicio_stats, column=2, value="ESTADÍSTICAS").font = Font(bold=True)
        ws.cell(row=fila_inicio_stats+1, column=2, value="Métrica").font = Font(bold=True)
        ws.cell(row=fila_inicio_stats+1, column=3, value=local_name).font = Font(bold=True)
        ws.cell(row=fila_inicio_stats+1, column=4, value=visitante_name).font = Font(bold=True)
        
        for i, stat in enumerate(estadisticas_partido):
            ws.cell(row=fila_inicio_stats + 2 + i, column=2, value=stat[0])
            ws.cell(row=fila_inicio_stats + 2 + i, column=3, value=stat[1])
            ws.cell(row=fila_inicio_stats + 2 + i, column=4, value=stat[2])

        ws.cell(row=fila_inicio_stats, column=5, value="PUNTOS ALTOS").font = Font(bold=True, color="008000")
        ws.cell(row=fila_inicio_stats, column=6, value="PUNTOS BAJOS").font = Font(bold=True, color="FF0000")
        for i, p in enumerate(puntos_altos): ws.cell(row=fila_inicio_stats + 1 + i, column=5, value=f"• {p}")
        for i, p in enumerate(puntos_bajos): ws.cell(row=fila_inicio_stats + 1 + i, column=6, value=f"• {p}")

        if os.path.exists(img_parado):
            img1 = XLImage(img_parado); img1.width, img1.height = 500, 360 
            ws.add_image(img1, f'H{fila_inicio_stats}') 
        if os.path.exists(img_shots):
            img2 = XLImage(img_shots); img2.width, img2.height = 500, 360
            ws.add_image(img2, f'P{fila_inicio_stats}') 
        if os.path.exists(img_shots_rival):
            img3 = XLImage(img_shots_rival); img3.width, img3.height = 500, 360
            ws.add_image(img3, f'H{fila_inicio_stats + 20}') 

        wb.save(RUTA_EXCEL)
        wb.close()
        
        for f in [img_parado, img_shots, img_shots_rival]:
            if os.path.exists(f): os.remove(f)
            
        print(f"✔️ Hoja '{nombre_hoja}' guardada con éxito.")

    except Exception as e:
        print(f"❌ Falló el procesamiento del partido {event_id}: {e}")

# =========================================================
# EJECUCIÓN DEL CRAWLER
# =========================================================
if __name__ == "__main__":
    print(f"Iniciando extracción masiva para {EQUIPO_OBJETIVO} (Temporada {ANIO_TEMPORADA})...")
    
    partidos = obtener_partidos_liga(TEAM_ID, ANIO_TEMPORADA)
    
    partidos_cronologicos = list(reversed(partidos))
    
    for partido in partidos_cronologicos:
        ejecutar_reporte_exacto(partido['id'], EQUIPO_OBJETIVO)
        time.sleep(3)
        
    print("\n✅ ¡PROCESO FINALIZADO CON ÉXITO! El archivo Excel está listo en data/equipos/.")